import asyncio
import io
import logging
from collections.abc import Sequence
from typing import Any

from celery.app.task import Task
from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import select

from src.celery_app import celery_app
from src.data.models.batch import Batch
from src.data.models.product import Product
from src.storage.minio_service import minio_service
from src.tasks.aggregation import AsyncSessionLocal

logger = logging.getLogger(__name__)


@celery_app.task(bind=True, max_retries=3)  # type: ignore[untyped-decorator]
def generate_batch_report(
    self: Task,
    batch_id: int,
    format: str = "excel",
    user_email: str | None = None,
) -> dict[str, Any]:
    try:
        return asyncio.run(
            _process_report_generation(batch_id, format, user_email)
        )
    except Exception as exc:
        logger.error(
            "Ошибка генерации отчёта для партии %s: %s", batch_id, exc
        )
        raise self.retry(exc=exc, countdown=60)


async def _process_report_generation(
    batch_id: int, report_format: str, user_email: str | None
) -> dict[str, Any]:

    async with AsyncSessionLocal() as session:
        batch = (
            await session.execute(select(Batch).where(Batch.id == batch_id))
        ).scalar_one_or_none()
        if not batch:
            raise ValueError(f"Партия с id {batch_id} не найдена")

        products = (
            (
                await session.execute(
                    select(Product)
                    .where(Product.batch_id == batch_id)
                    .order_by(Product.id)
                )
            )
            .scalars()
            .all()
        )

    if report_format == "excel":
        file_content, file_name = _generate_excel_report(batch, products)
        content_type = (
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
    else:
        raise ValueError("Пока поддерживается только формат 'excel'")

    file_url, expires_at = minio_service.upload_file_bytes(
        bucket="reports",
        file_bytes=file_content,
        object_name=file_name,
        content_type=content_type,
        expires_days=7,
    )
    if user_email:
        logger.info(
            "Отчёт готов, уведомление будет отправлено на %s", user_email
        )

    return {
        "success": True,
        "file_url": file_url,
        "file_name": file_name,
        "file_size": len(file_content),
        "expires_at": expires_at,
    }


def _generate_excel_report(
    batch: Batch, products: Sequence[Product]
) -> tuple[bytes, str]:
    """
    Генерация Excel-файла
    """
    wb = Workbook()

    ws_info = wb.active
    ws_info.title = "Информация о партии"

    info_data = [
        ("Номер партии", str(batch.batch_number)),
        ("Дата партии", str(batch.batch_date)),
        ("Статус", "Закрыта" if batch.is_closed else "Открыта"),
        (
            "Рабочий центр",
            (
                str(batch.work_center.identifier)
                if batch.work_center
                else "Не указан"
            ),
        ),
        ("Смена", batch.shift),
        ("Бригада", batch.team),
        ("Номенклатура", batch.nomenclature),
        (
            "Начало смены",
            (
                batch.shift_start.strftime("%Y-%m-%d %H:%M:%S")
                if batch.shift_start
                else "-"
            ),
        ),
        (
            "Окончание смены",
            (
                batch.shift_end.strftime("%Y-%m-%d %H:%M:%S")
                if batch.shift_end
                else "-"
            ),
        ),
    ]

    for row_idx, (key, value) in enumerate(info_data, start=1):
        ws_info.cell(row=row_idx, column=1, value=key).font = Font(bold=True)
        ws_info.cell(row=row_idx, column=2, value=value)

    ws_info.column_dimensions["A"].width = 20
    ws_info.column_dimensions["B"].width = 30

    ws_products = wb.create_sheet("Продукция")
    headers_products = [
        "ID",
        "Уникальный код",
        "Аггрегирована",
        "Дата аггрегации",
    ]
    ws_products.append(headers_products)

    for p in products:
        agg_status = "Да" if p.is_aggregated else "Нет"
        agg_date = (
            p.aggregated_at.strftime("%Y-%m-%d %H:%M:%S")
            if p.aggregated_at
            else "-"
        )
        ws_products.append([p.id, p.unique_code, agg_status, agg_date])

    ws_stats = wb.create_sheet("Статистика")

    total = len(products)
    aggregated = sum(1 for p in products if p.is_aggregated)
    remaining = total - aggregated
    percent = (aggregated / total * 100) if total > 0 else 0.0

    shift_hours = 8
    # shift_hours = (batch.shift_end - batch.shift_start).total_seconds() / 3600
    speed = round(aggregated / shift_hours, 1) if shift_hours > 0 else 0.0

    stats_data: list[tuple[str, str]] = [
        ("Всего продукции", str(total)),
        ("Аггрегировано", str(aggregated)),
        ("Осталось", str(remaining)),
        ("Процент выполнения", f"{percent:.0f}%"),
        ("Средняя скорость", f"{speed} ед/час"),
    ]

    for row_idx, (key, value) in enumerate(stats_data, start=1):
        ws_stats.cell(row=row_idx, column=1, value=key).font = Font(bold=True)
        ws_stats.cell(row=row_idx, column=2, value=value)

    ws_stats.column_dimensions["A"].width = 25
    ws_stats.column_dimensions["B"].width = 20

    file_stream = io.BytesIO()
    wb.save(file_stream)
    return file_stream.getvalue(), f"batch_{batch.id}_report.xlsx"
